"""Read MCP_SETUP_GUIDE.md and write MCP_SETUP_GUIDE.docx."""
import pathlib
import re

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor


INPUT_MD = pathlib.Path(__file__).parent / "MCP_SETUP_GUIDE.md"
OUTPUT_DOCX = pathlib.Path(__file__).parent / "MCP_SETUP_GUIDE.docx"


def _add_code_shading(paragraph):
    """Apply light-gray background shading to a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)


def _set_code_font(run):
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def build_docx(md_text: str) -> Document:
    doc = Document()

    in_fence = False
    fence_lines: list[str] = []

    def flush_fence():
        nonlocal fence_lines
        for line in fence_lines:
            p = doc.add_paragraph(style="Normal")
            _add_code_shading(p)
            run = p.add_run(line)
            _set_code_font(run)
        fence_lines = []

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()

        # code fence open/close
        if line.startswith("```"):
            if not in_fence:
                in_fence = True
            else:
                in_fence = False
                flush_fence()
            continue

        if in_fence:
            fence_lines.append(line)
            continue

        # headings
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        # horizontal rule
        elif line.startswith("---"):
            doc.add_paragraph()
        # bullet list
        elif re.match(r"^[-*] ", line):
            doc.add_paragraph(line[2:], style="List Bullet")
        # numbered list
        elif re.match(r"^\d+\. ", line):
            text = re.sub(r"^\d+\. ", "", line)
            doc.add_paragraph(text, style="List Number")
        # table row — convert to plain text, skip separator rows
        elif line.startswith("|"):
            if re.match(r"^\|[-| ]+\|$", line):
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            doc.add_paragraph("  |  ".join(cells), style="Normal")
        # blank line
        elif line == "":
            doc.add_paragraph()
        else:
            # strip simple inline markdown: **bold**, *italic*, `code`
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            clean = re.sub(r"\*(.+?)\*", r"\1", clean)
            clean = re.sub(r"`(.+?)`", r"\1", clean)
            doc.add_paragraph(clean, style="Normal")

    return doc


if __name__ == "__main__":
    md_text = INPUT_MD.read_text(encoding="utf-8")
    doc = build_docx(md_text)
    doc.save(OUTPUT_DOCX)
    print(f"Written: {OUTPUT_DOCX}")
